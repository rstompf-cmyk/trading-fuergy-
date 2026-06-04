# -*- coding: utf-8 -*-
"""report.py – export plánu D-1 do Excelu (hárky Plán + Súhrn + graf)."""
from __future__ import annotations
import datetime as dt
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, numbers
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")


def build_plan_excel(schedule, summary, meta, path):
    wb = Workbook()
    ws = wb.active; ws.title = "Plán"

    ws["A1"] = f"Plán D-1 obchodovania FTV + batéria – {meta.get('date','')}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"Lokalita {meta.get('lat')},{meta.get('lon')}  •  FTV {meta.get('kwp')} kWp  •  "
                f"batéria {meta.get('batt_kw')} kW / {meta.get('batt_kwh')} kWh  •  "
                f"nabíjanie zo siete: {'áno' if meta.get('allow_grid_charge') else 'nie'}")
    ws["A2"].font = Font(italic=True, color="808080")

    cols = [("hod", "hour"), ("FTV [kWh]", "pv_kwh"), ("ISOT [€/MWh]", "price_eur"),
            ("Batéria [kW] (+vyb/−nab)", "batt_kw"), ("Sieť [kWh] (+pred/−nák)", "grid_kwh"),
            ("Obchod [MWh]", "order_mwh"), ("Orezané [kWh]", "curtail_kwh"), ("SOC [%]", "soc_pct")]
    hr = 4
    for j, (label, _) in enumerate(cols, start=1):
        c = ws.cell(hr, j, label); c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, (_, row) in enumerate(schedule.iterrows(), start=hr+1):
        for j, (_, key) in enumerate(cols, start=1):
            v = row[key]
            cell = ws.cell(i, j, float(v) if key != "hour" else int(v))
            if key in ("pv_kwh", "price_eur", "batt_kw", "grid_kwh", "curtail_kwh"):
                cell.number_format = "0.0"
            if key == "order_mwh":
                cell.number_format = "0.000"
            if key == "soc_pct":
                cell.number_format = "0"
    last = hr + len(schedule)
    for j in range(1, len(cols)+1):
        ws.column_dimensions[get_column_letter(j)].width = 15 if j > 1 else 6

    # graf: ISOT + batéria (ľavá os), SOC % (pravá os)
    ch = LineChart(); ch.title = "Cena, batéria a SOC počas dňa"
    ch.height, ch.width = 8, 20
    ch.y_axis.title = "€/MWh, kW"
    data = Reference(ws, min_col=3, max_col=4, min_row=hr, max_row=last)
    cats = Reference(ws, min_col=1, min_row=hr+1, max_row=last)
    ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
    soc = LineChart()
    soc.add_data(Reference(ws, min_col=8, max_col=8, min_row=hr, max_row=last), titles_from_data=True)
    soc.y_axis.axId = 200; soc.y_axis.title = "SOC %"
    soc.y_axis.crosses = "max"
    ch += soc
    ws.add_chart(ch, f"A{last+3}")

    # hárok Súhrn
    s = wb.create_sheet("Súhrn")
    s["A1"] = "Súhrn plánu"; s["A1"].font = TITLE_FONT
    labels = {
        "ZISK_EUR": "ZISK [€]", "trzba_export_EUR": "Tržba z predaja [€]",
        "naklad_import_EUR": "Náklad na nákup [€]", "naklad_cyklus_EUR": "Náklad na cyklus [€]",
        "bez_baterie_EUR": "Bez batérie [€]", "prinos_baterie_EUR": "Prínos batérie [€]",
        "nabite_kWh": "Nabité [kWh]", "vybite_kWh": "Vybité [kWh]",
        "import_kWh": "Nákup zo siete [kWh]", "orezane_kWh": "Orezaná FTV [kWh]"}
    r = 3
    for k, lab in labels.items():
        s.cell(r, 1, lab).font = Font(bold=(k == "ZISK_EUR"))
        c = s.cell(r, 2, summary.get(k)); c.number_format = "0.00"
        if k == "ZISK_EUR":
            c.font = Font(bold=True, size=12, color="2E7D32")
        r += 1
    r += 1
    s.cell(r, 1, "Parametre:").font = Font(bold=True); r += 1
    for k in ("date", "lat", "lon", "kwp", "tilt", "azimuth", "eff", "batt_kw", "batt_kwh",
              "eff_c", "eff_d", "soc_min_pct", "soc_max_pct", "soc_init_pct",
              "grid_kw", "grid_fee", "cycle_cost", "allow_grid_charge", "terminal_soc_pct",
              "min_spread_eur", "min_trade_mwh", "price_scale", "pv_scale", "block_neg_import"):
        if k in meta:
            s.cell(r, 1, k); s.cell(r, 2, str(meta[k])); r += 1
    s.cell(r+1, 1, f"Vygenerované {dt.datetime.now():%Y-%m-%d %H:%M}").font = Font(italic=True, color="808080")
    s.column_dimensions["A"].width = 24; s.column_dimensions["B"].width = 22

    wb.save(path)
    return path
