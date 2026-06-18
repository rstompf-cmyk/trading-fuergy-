# -*- coding: utf-8 -*-
"""
tools/export_livesim_xlsx.py — CLI export livesim tabuľky do Excelu (15-min + hodinový priemer).

Rovnaký výstup ako tlačidlo /livesim_table_xlsx, ale spustiteľné z príkazu (napr. v kontajneri
na Windows) pre ľubovoľný profil / case / deň / rozsah — bez UI.

Použitie (v kontajneri):
    docker exec trading-fuergy-dev python tools/export_livesim_xlsx.py \
        --case dt_15min --day 2026-06-12 --out /app/out/export.xlsx
    # rozsah:
    docker exec trading-fuergy-dev python tools/export_livesim_xlsx.py \
        --case dt_15min --from 2026-06-01 --to 2026-06-18 --out /app/out/export.xlsx

Potom skopíruj von:  docker cp trading-fuergy-dev:/app/out/export.xlsx .
(--port default = env PORT/APP_PORT/8000; --profile vynúti profil cez FTV_PROFILE.)
"""
from __future__ import annotations
import argparse, os, sys
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

HEADERS = ["slot_start", "slot_end", "batt_kw_net", "batt_dam_kw", "batt_vdt_kw",
           "work_kwh", "soc_pct", "ftv_kw", "dt_eur_mwh", "dt_real_eur_mwh", "vdt_eur_mwh"]


def _agg_15(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["_slot"] = pd.to_datetime(df["time"]).dt.floor("15min")
    agg = {"soc_pct": "last", "ftv_kw": "mean", "dt_eur": "mean"}
    for c in ("plan_batt_kw", "plan_batt_dam_kw", "plan_batt_vdt_kw", "dt_real_eur", "vdt_eur"):
        if c in df.columns:
            agg[c] = "mean"
    g = df.groupby("_slot").agg(agg).reset_index()
    g["work_kwh"] = g.get("plan_batt_kw", 0) * 0.25
    return g


def _to_hourly(g15: pd.DataFrame) -> pd.DataFrame:
    gh = g15.copy()
    gh["_slot"] = pd.to_datetime(gh["_slot"]).dt.floor("1h")
    aggh = {k: ("sum" if k == "work_kwh" else ("last" if k == "soc_pct" else "mean"))
            for k in g15.columns if k != "_slot"}
    return gh.groupby("_slot").agg(aggh).reset_index()


def _rows(gg: pd.DataFrame, freq: str):
    out = []
    for _, r in gg.iterrows():
        ts = pd.Timestamp(r["_slot"])
        te = ts + (pd.Timedelta(minutes=15) if freq == "15" else pd.Timedelta(hours=1))
        out.append([ts.strftime("%Y-%m-%d %H:%M"), te.strftime("%H:%M"),
                    round(float(r.get("plan_batt_kw", 0)), 1),
                    round(float(r.get("plan_batt_dam_kw", 0)), 1),
                    round(float(r.get("plan_batt_vdt_kw", 0)), 1),
                    round(float(r.get("work_kwh", 0)), 2),
                    round(float(r.get("soc_pct", 0)), 1),
                    round(float(r.get("ftv_kw", 0)), 1),
                    round(float(r.get("dt_eur", 0)), 1),
                    round(float(r.get("dt_real_eur", 0)), 1),
                    round(float(r.get("vdt_eur", 0)), 1)])
    return out


def _write_sheet(wb, title, rows, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    if first:
        ws.title = title
    ws.append(HEADERS)
    hf = PatternFill("solid", fgColor="1F4E78"); ff = Font(bold=True, color="FFFFFF")
    for c in ws[1]:
        c.fill = hf; c.font = ff; c.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--case", default="dt_15min")
    ap.add_argument("--day")
    ap.add_argument("--from", dest="dfrom")
    ap.add_argument("--to", dest="dto")
    ap.add_argument("--port", default=os.environ.get("PORT", os.environ.get("APP_PORT", "8000")))
    ap.add_argument("--profile", default=None, help="vynúti profil cez FTV_PROFILE")
    ap.add_argument("--out", default="out/livesim_export.xlsx")
    a = ap.parse_args()
    if a.profile:
        os.environ["FTV_PROFILE"] = a.profile
    import livesim as lsim

    if a.day:
        days = [a.day]
    elif a.dfrom and a.dto:
        days = [d.date().isoformat() for d in pd.date_range(a.dfrom, a.dto, freq="D")]
    else:
        days = [pd.Timestamp.today().date().isoformat()]

    # Fallback: ak load_series netrafí (market/profile env), načítaj CSV priamo z disku.
    _raw_cache = {}

    def _raw_df():
        if "df" in _raw_cache:
            return _raw_cache["df"]
        import glob
        pats = [f"out/**/livesim_{a.case}*{a.port}*.csv", f"out/**/livesim_{a.case}*.csv",
                f"out/livesim_{a.case}*{a.port}*.csv"]
        hits = []
        for p in pats:
            hits += glob.glob(p, recursive=True)
        df = None
        for f in sorted(set(hits)):
            try:
                t = pd.read_csv(f, parse_dates=["time"])
                if "time" in t.columns and not t.empty:
                    df = t
                    print(f"  (fallback: {f})")
                    break
            except Exception:
                continue
        _raw_cache["df"] = df
        return df

    parts = []
    for d in days:
        df = None
        try:
            df = lsim.load_series(a.case, port=str(a.port), day=d, max_points=10**9)
        except Exception as e:
            print(f"  {d}: load_series zlyhal: {e}")
        if df is None or df.empty:
            raw = _raw_df()
            if raw is not None:
                df = raw[pd.to_datetime(raw["time"]).dt.date.astype(str) == d].copy()
        if df is None or df.empty:
            print(f"  {d}: žiadne dáta (case={a.case}, port={a.port})")
            continue
        parts.append(_agg_15(df))
    if not parts:
        print("Žiadne dáta — skontroluj --case/--port/--profile a či livesim zbehol.")
        sys.exit(2)
    g15 = pd.concat(parts, ignore_index=True).sort_values("_slot")
    gh = _to_hourly(g15)
    wb = openpyxl.Workbook()
    _write_sheet(wb, "15-min", _rows(g15, "15"), first=True)
    _write_sheet(wb, "Hodinový (priemer)", _rows(gh, "h"))
    os.makedirs(os.path.dirname(a.out) or ".", exist_ok=True)
    wb.save(a.out)
    print(f"OK: {len(g15)} × 15-min + {len(gh)} hodinových riadkov → {a.out}")


if __name__ == "__main__":
    main()
